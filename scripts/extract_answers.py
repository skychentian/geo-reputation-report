#!/usr/bin/env python3
"""Extract a GEO diagnosis analysis package into a normalized answers.json.

Input: a package directory containing diagnosis-answers.xlsx + manifest.json,
or the xlsx path directly.

The script does not judge sentiment. It only normalizes dimensions, maps
question topics to report layers, and selects the target/competitor brands.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")


PLATFORM_ALIASES = {
    "豆包手机版": "豆包", "豆包": "豆包", "doubao": "豆包",
    "deepseek手机版": "DeepSeek", "deepseek": "DeepSeek", "DeepSeek": "DeepSeek",
    "元宝手机版": "腾讯元宝", "腾讯元宝": "腾讯元宝", "元宝": "腾讯元宝",
    "通义千问手机版": "通义千问", "通义千问": "通义千问", "通义": "通义千问",
    "文心一言": "文心一言", "文心": "文心一言",
    "kimi": "Kimi", "Kimi": "Kimi",
}
GENERIC_COMPETITORS = {"橄榄油", "花生油", "菜籽油", "玉米油", "调和油", "大豆油", "粮油", "食用油"}
COLORS = ["#14684b", "#c08a2e", "#b45a3c", "#4f6d9a", "#7a5c9e", "#2f7f8f"]


def normalize_platform(raw: str) -> str:
    name = str(raw or "").strip()
    if name in PLATFORM_ALIASES:
        return PLATFORM_ALIASES[name]
    for key, value in PLATFORM_ALIASES.items():
        if name.lower() == key.lower():
            return value
    return re.sub(r"(手机版|网页版|App版)$", "", name) or "未知平台"


def infer_target(manifest: dict, override: str | None) -> str:
    if override:
        return override.strip()
    task = manifest.get("task") or {}
    raw = str(task.get("brandName") or "").strip()
    if not raw:
        title = str(task.get("title") or "").strip()
        raw = re.split(r"\s*(GEO|AI|诊断|报告)", title)[0] or title
    raw = re.sub(r"(美誉度|口碑|GEO诊断|诊断|报告)$", "", raw).strip()
    return raw or "目标品牌"


def pick_competitors(summary_rows: list[list], target: str, requested: list[str], limit: int) -> list[str]:
    if requested:
        return [x.strip() for x in requested if x.strip() and x.strip() != target]
    counts = []
    for row in summary_rows:
        if not row or len(row) < 3:
            continue
        name = str(row[1] or "").strip()
        if not name or name == target or name in GENERIC_COMPETITORS:
            continue
        try:
            mentions = int(float(row[2] or 0))
        except (TypeError, ValueError):
            mentions = 0
        if mentions > 0:
            counts.append((mentions, name))
    counts.sort(reverse=True)
    return [name for _, name in counts[:limit]]


def load_manifest(package: Path, xlsx: Path) -> dict:
    for candidate in (package / "manifest.json", xlsx.parent / "manifest.json"):
        if candidate.exists():
            return json.loads(candidate.read_text(encoding="utf-8"))
    return {}


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract diagnosis-answers.xlsx into answers.json")
    parser.add_argument("package", help="package directory or diagnosis-answers.xlsx")
    parser.add_argument("--out", default="work/answers.json")
    parser.add_argument("--target-brand", default=None)
    parser.add_argument("--competitors", default=None, help="comma separated; default top 2 from competitor_summary")
    parser.add_argument("--competitor-limit", type=int, default=2)
    parser.add_argument("--layer-map", default=str(Path(__file__).resolve().parent.parent / "config" / "layer-map.json"))
    parser.add_argument("--default-layer", default="品牌整体")
    args = parser.parse_args()

    pkg = Path(args.package).expanduser().resolve()
    xlsx = pkg / "diagnosis-answers.xlsx" if pkg.is_dir() else pkg
    if not xlsx.exists():
        sys.exit(f"xlsx not found: {xlsx}")

    manifest = load_manifest(pkg, xlsx)
    layer_map = json.loads(Path(args.layer_map).read_text(encoding="utf-8"))
    topic_to_layer = {topic: layer for layer, topics in layer_map["layers"].items() for topic in topics}
    layers = list(layer_map["layers"].keys())

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "answers" not in wb.sheetnames:
        sys.exit("xlsx is missing the answers sheet")
    ws = wb["answers"]
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(headers)}
    required = ["index", "question", "stage", "topic", "platform_name", "round", "status", "answer_content"]
    missing = [name for name in required if name not in idx]
    if missing:
        sys.exit(f"answers sheet missing columns: {missing}")

    target = infer_target(manifest, args.target_brand)
    summary_rows = []
    if "competitor_summary" in wb.sheetnames:
        summary_rows = list(wb["competitor_summary"].iter_rows(min_row=2, values_only=True))
    requested = args.competitors.split(",") if args.competitors else []
    competitors = pick_competitors(summary_rows, target, requested, args.competitor_limit)
    if not competitors:
        sys.exit("no competitors found; pass --competitors explicitly")

    completed = str((manifest.get("task") or {}).get("completedAt") or "")
    if completed:
        period = completed[:10]
    else:
        title = str((manifest.get("task") or {}).get("title") or "")
        match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", title)
        period = f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}" if match else date.today().isoformat()

    answers = []
    unmapped = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[idx["index"]] is None:
            continue
        topic = str(row[idx["topic"]] or "").strip()
        layer = topic_to_layer.get(topic)
        if layer is None:
            unmapped[topic] += 1
            layer = args.default_layer
        platform_raw = str(row[idx["platform_name"]] or "").strip()
        text = str(row[idx["answer_content"]] or "")
        answers.append({
            "id": f"{period}-{row[idx['index']]}",
            "index": row[idx["index"]],
            "question": str(row[idx["question"]] or "").strip(),
            "stage": str(row[idx["stage"]] or "").strip(),
            "topic": topic,
            "layer": layer,
            "platform": normalize_platform(platform_raw),
            "platformRaw": platform_raw,
            "round": row[idx["round"]],
            "text": text,
            "charCount": len(text),
        })

    if not answers:
        sys.exit("no answers extracted")

    platform_counts = Counter(a["platform"] for a in answers)
    manifest_platforms = [normalize_platform(x) for x in ((manifest.get("task") or {}).get("platforms") or [])]
    platform_order = []
    for platform in manifest_platforms + [p for p, _ in platform_counts.most_common()]:
        if platform in platform_counts and platform not in platform_order:
            platform_order.append(platform)

    brand_line = " · ".join([target] + competitors)
    result = {
        "meta": {
            "package": str(pkg),
            "taskId": (manifest.get("task") or {}).get("id"),
            "reportTitle": (manifest.get("task") or {}).get("title") or f"{target}美誉度",
            "targetBrand": target,
            "competitors": competitors,
            "brandLine": brand_line,
            "period": period,
            "periodLabel": f"{period[5:7]}/{period[8:10]}",
            "completedAt": completed or None,
            "questionCount": len({a["question"] for a in answers}),
            "answerCount": len(answers),
            "platforms": platform_order,
            "platformCounts": dict(platform_counts),
            "layers": layers,
            "layerCounts": dict(Counter(a["layer"] for a in answers)),
            "unmappedTopics": dict(unmapped),
        },
        "answers": answers,
    }

    out = Path(args.out).expanduser().resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "out": str(out),
        "target": target,
        "competitors": competitors,
        "answers": len(answers),
        "questions": result["meta"]["questionCount"],
        "layers": result["meta"]["layerCounts"],
        "platforms": platform_order,
        "unmappedTopics": dict(unmapped),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
